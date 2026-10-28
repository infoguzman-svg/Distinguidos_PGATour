"""Convierte Selecciones_AAAA.xlsx en data/datos.json (lo que lee la página)."""
import datetime as dt, json, re, sys, unicodedata
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
archivos = sorted(RAIZ.glob("Selecciones_*.xlsx"))
if not archivos:
    sys.exit("No encuentro Selecciones_AAAA.xlsx en la carpeta del proyecto")
XLSX = archivos[-1]
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

def sin_tildes(s):
    s = unicodedata.normalize("NFD", str(s))
    return "".join(c for c in s if unicodedata.category(c) != "Mn")

def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", sin_tildes(s).lower()).strip("-")

def txt(v):
    return str(v).strip() if v is not None else ""

def fecha(v):
    if isinstance(v, dt.datetime): return v.date().isoformat()
    if isinstance(v, dt.date): return v.isoformat()
    s = txt(v)
    for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try: return dt.datetime.strptime(s, f).date().isoformat()
        except ValueError: pass
    if s: print(f"ADVERTENCIA: fecha no reconocida: {s}")
    return None

wb = openpyxl.load_workbook(XLSX, data_only=True)
def filas(hoja):
    if hoja not in wb.sheetnames: return
    for r in wb[hoja].iter_rows(min_row=2, values_only=True):
        if any(c is not None and txt(c) for c in r):
            yield list(r) + [None] * 8

participantes = [txt(r[0]) for r in filas("Participantes") if txt(r[0])]
escala = [[int(r[0]), int(r[1]), int(r[2])] for r in filas("Puntos") if isinstance(r[0], (int, float)) and isinstance(r[1], (int, float))]
clases = {txt(r[0]): [int(r[1] or 0), int(r[2] or 0), int(r[3] or 0)] for r in filas("Clases") if txt(r[0]) and isinstance(r[1], (int, float))}

torneos, ids, actual = [], set(), None
for r in filas("Torneos"):
    nombre, clase = txt(r[0]), txt(r[1])
    if not nombre: continue
    if clase not in clases:
        print(f"ADVERTENCIA: clase '{clase}' del torneo '{nombre}' no existe en la hoja Clases")
    base = slug(nombre) or "torneo"; id_, n = base, 2
    while id_ in ids: id_, n = f"{base}-{n}", n + 1
    ids.add(id_)
    if txt(r[4]).lower() in ("sí", "si", "x"): actual = id_
    ini = fecha(r[2])
    mes = txt(r[5]) or (MESES[int(ini[5:7]) - 1] if ini else "")
    torneos.append({"id": id_, "n": nombre, "c": clase, "ini": ini, "fin": fecha(r[3]), "m": mes, "sel": {}})
por_nombre = {t["n"]: t for t in torneos}

for r in filas("Selecciones"):
    nombre, part = txt(r[0]), txt(r[1])
    if nombre not in por_nombre:
        print(f"ADVERTENCIA: torneo '{nombre}' de Selecciones no está en la hoja Torneos"); continue
    if part not in participantes:
        print(f"ADVERTENCIA: participante '{part}' no está en la hoja Participantes")
    jug = [txt(c) for c in r[2:6] if txt(c)]
    if len(jug) != 4:
        print(f"ADVERTENCIA: {part} en '{nombre}' tiene {len(jug)} jugadores (se esperan 4)")
    por_nombre[nombre]["sel"][part] = jug

for r in filas("Puestos manuales"):
    nombre, part = txt(r[0]), txt(r[1])
    if nombre not in por_nombre or part not in participantes:
        print(f"ADVERTENCIA: fila de Puestos manuales no reconocida: {nombre} / {part}"); continue
    t = por_nombre[nombre]
    t.setdefault("puestos", [0] * len(participantes))
    t["puestos"][participantes.index(part)] = int(r[2] or 0)

temporada = (re.search(r"(\d{4})", XLSX.stem) or [None, str(dt.date.today().year)])[1]
datos = {"temporada": temporada, "participantes": participantes, "escalaPos": escala, "clases": clases, "actual": actual, "torneos": torneos}
(RAIZ / "data").mkdir(exist_ok=True)
(RAIZ / "data" / "datos.json").write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"datos.json listo: {len(torneos)} torneos, {len(participantes)} participantes ({XLSX.name})")
